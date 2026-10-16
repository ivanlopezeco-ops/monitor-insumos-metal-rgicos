"""Exporta config/catalogo.yaml a una planilla para revisión (results/catalogo_revision.xlsx)."""
import pathlib
import sys

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = pathlib.Path(__file__).resolve().parents[1]
cat = yaml.safe_load((BASE / "config" / "catalogo.yaml").read_text(encoding="utf-8"))
salida = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else BASE / "results" / "catalogo_revision.xlsx"

TIPO = {"precio_real": "Precio real", "futuro": "Futuro", "precio_comercio_exterior": "Precio de comercio exterior",
        "indice": "Índice", "estimacion": "Estimación"}
LIC = {"uso_interno": "Solo uso interno", "publicable_citando": "Publicable citando", "publicable_con_aviso": "Publicable con crédito y aviso",
       "a_confirmar": "A confirmar"}
EST = {"automatico": "Automático", "acumula": "Automático, se acumula", "semiautomatico": "Semiautomático", "pendiente": "Pendiente"}
AZUL = PatternFill("solid", fgColor="1F4E79")


def hoja(wb, titulo, encabezados, filas, anchos):
    ws = wb.create_sheet(titulo)
    ws.append(encabezados)
    for f in filas:
        ws.append(f)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = AZUL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = "D2" if len(encabezados) > 6 else "A2"
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()
wb.remove(wb.active)
filas = []
for ins in cat["insumos"]:
    for s in ins["series"]:
        f = s.get("fuente", {})
        filas.append([ins["grupo"], ins["nombre"], s["nombre"], s.get("mercado", ""), TIPO.get(s.get("tipo"), s.get("tipo")),
                      s.get("que_mide", ""), s.get("limitaciones", ""), f.get("nombre", ""), f.get("url", ""),
                      s.get("frecuencia", ""), s.get("rezago", ""), s.get("desde", ""), s.get("unidad_original", ""),
                      s.get("unidad_tablero", ""), s.get("conversion", ""), LIC.get(s.get("licencia"), s.get("licencia")),
                      EST.get(s.get("estado"), s.get("estado")), "", s["id"]])
hoja(wb, "Series", ["Grupo", "Insumo", "Serie", "Mercado", "Tipo de dato", "Qué mide", "Limitaciones", "Fuente", "Link",
                    "Frecuencia", "Rezago", "Desde", "Unidad original", "Unidad en el tablero", "Conversión", "Licencia",
                    "Estado", "Comentarios de Iván", "id"],
     filas, [16, 22, 38, 11, 16, 60, 45, 26, 30, 11, 14, 16, 16, 14, 26, 18, 16, 30, 24])
hoja(wb, "Insumos", ["Grupo", "Insumo", "Qué es", "Por qué importa", "Cantidad de series", "Comentarios de Iván"],
     [[i["grupo"], i["nombre"], i["que_es"], i["por_que_importa"], len(i["series"]), ""] for i in cat["insumos"]],
     [18, 28, 60, 60, 10, 30])
hoja(wb, "Tipos de dato", ["Tipo", "Explicación"], [[TIPO[k], v] for k, v in cat["tipos"].items()], [26, 100])
salida.parent.mkdir(exist_ok=True, parents=True)
wb.save(salida)
print(f"{salida}: {len(filas)} series en {len(cat['insumos'])} insumos")
